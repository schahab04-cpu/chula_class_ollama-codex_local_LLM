import json
import shutil
import unittest
from pathlib import Path
from uuid import uuid4
from unittest.mock import patch
from docx import Document
from pypdf import PdfWriter
from src.rag import KnowledgeBase, extract, DocumentError, embed, split_passage
from src.diagnose import audit


class RagTests(unittest.TestCase):
    def setUp(self):
        self.folder = Path(__file__).resolve().parent / ("rag-" + uuid4().hex)
        self.folder.mkdir()
        self.addCleanup(shutil.rmtree, self.folder)
        self.path = self.folder / "abc--notes.txt"
        self.path.write_text("The launch date is 15 June 2027.", encoding="utf-8")
        self.sources = [{"file_id": self.path.name, "name": "notes.txt", "location": "Paragraph 1",
                         "text": "The launch date is 15 June 2027."}]

    def test_text_markdown_and_docx_extraction(self):
        self.assertIn("15 June", extract(self.path)[0]["text"])
        md = self.folder / "notes.md"
        md.write_text("# Launch\n\nThe launch date is 15 June 2027.", encoding="utf-8")
        self.assertEqual(extract(md)[1]["location"], "Paragraph 2")
        doc = Document()
        doc.add_paragraph("Launch date")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "15 June 2027"
        path = self.folder / "notes.docx"
        doc.save(path)
        self.assertIn("15 June", extract(path)[1]["text"])

    def test_invalid_empty_binary_and_scanned_files(self):
        for name, data in (("x.txt", b""), ("x.md", b"a\x00b"), ("x.docx", b"bad"),
                           ("x.pdf", b"bad"), ("x.exe", b"data")):
            path = self.folder / name
            path.write_bytes(data)
            with self.assertRaises(DocumentError):
                extract(path)
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        path = self.folder / "scan.pdf"
        with path.open("wb") as target:
            writer.write(target)
        with self.assertRaisesRegex(DocumentError, "OCR"):
            extract(path)

    @patch("src.rag.embed", return_value=[[1.0, 0.0]])
    def test_persistent_index_search_and_changed_file(self, embeddings):
        kb = KnowledgeBase(self.folder)
        kb.index(self.path)
        restarted = KnowledgeBase(self.folder)
        self.assertTrue(restarted.ready(self.path))
        self.assertEqual(restarted.retrieve("When is launch?")[0]["name"], "notes.txt")
        self.path.write_text("Changed", encoding="utf-8")
        self.assertFalse(restarted.ready(self.path))
        with self.assertRaises(DocumentError):
            restarted.retrieve("When is launch?")

    @patch("src.rag.embed", return_value=[[1.0, 0.0]])
    def test_deleted_files_are_not_retrieved(self, embeddings):
        kb = KnowledgeBase(self.folder)
        kb.index(self.path)
        self.path.unlink()
        self.assertEqual(kb.retrieve("launch"), [])


    @patch("src.ollama_client.ollama", return_value={"embeddings": [[1.0, 0.0]]})
    def test_embedding_prefixes_and_no_truncation(self, ollama):
        embed(["question"], "nomic-embed-text:latest", query=True)
        self.assertEqual(ollama.call_args.args[1]["input"], ["search_query: question"])
        self.assertFalse(ollama.call_args.args[1]["truncate"])


    def test_long_chunks_preserve_words_and_sentence_boundaries(self):
        text = " ".join(f"Sentence {i} contains useful lecture information." for i in range(90))
        chunks = list(split_passage(text))
        self.assertTrue(all(len(c) <= 1000 for c in chunks))
        self.assertTrue(all(c.endswith(".") for c in chunks))
        for i in range(90):
            self.assertTrue(any(f"Sentence {i} contains useful lecture information." in c for c in chunks))

    @patch("src.rag.embed")
    def test_lexical_hit_survives_low_cosine_and_semantic_distractors(self, embeddings):
        kb = KnowledgeBase(self.folder)
        embeddings.return_value = [[0.1, (0.99 ** 0.5)]]
        kb.index(self.path)
        for i in range(12):
            path = self.folder / f"{i}--noise.txt"
            path.write_text("Unrelated weather observations.", encoding="utf-8")
            embeddings.return_value = [[1.0, 0.0]]
            kb.index(path)
        sources = kb.retrieve("launch date")
        self.assertLessEqual(len(sources), 8)
        self.assertIn(self.path.name, [s["file_id"] for s in sources])
        hit = next(s for s in sources if s["file_id"] == self.path.name)
        self.assertAlmostEqual(hit["score"], 0.1)

    @patch("src.rag.embed", return_value=[[1.0, 0.0]])
    def test_failed_embedding_keeps_previous_complete_index(self, embeddings):
        kb = KnowledgeBase(self.folder)
        kb.index(self.path)
        self.path.write_text("New text", encoding="utf-8")
        embeddings.side_effect = ValueError("Embedding failed")
        with self.assertRaises(ValueError):
            kb.index(self.path)
        with kb.connect() as db:
            stored = json.loads(db.execute("SELECT chunks FROM documents").fetchone()[0])
        self.assertIn("15 June", stored[0]["text"])
        self.assertFalse(kb.ready(self.path))

    @patch("src.rag.embed", return_value=[[1.0, 0.0]])
    def test_audit_checks_storage_without_reindexing(self, embeddings):
        kb = KnowledgeBase(self.folder)
        kb.index(self.path)
        before = (self.folder / "knowledge.sqlite3").read_bytes()
        report = audit(self.folder)[0]
        self.assertEqual(report["stored_chunks"], 1)
        self.assertEqual(report["vector_dimensions"], [2])
        self.assertTrue(report["fingerprint_matches"])
        self.assertTrue(report["extraction_matches_index"])
        self.assertEqual(before, (self.folder / "knowledge.sqlite3").read_bytes())

    @patch("src.rag.embed", return_value=[[1.0, 0.0]])
    def test_empty_or_corrupted_index_is_an_operational_error(self, embeddings):
        kb = KnowledgeBase(self.folder)
        kb.index(self.path)
        for stored in ([], [{"text": "valid text", "location": "Paragraph 1", "vector": [0, 0]}]):
            with kb.connect() as db:
                db.execute("UPDATE documents SET chunks=?", (json.dumps(stored),))
            with self.assertRaises((DocumentError, ValueError)):
                kb.retrieve("launch date")
